from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path
import os

BASE = Path(__file__).resolve().parents[1]
OUT = BASE / 'downloads'
OUT.mkdir(parents=True, exist_ok=True)

RED = 'C00000'
LIGHT_GRAY = 'F2F2F2'
MID_GRAY = 'D9E2F3'


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)


def set_no_split(row):
    trPr = row._tr.get_or_add_trPr()
    cantSplit = OxmlElement('w:cantSplit')
    trPr.append(cantSplit)


def set_run_font(run, east='FangSong', latin='Times New Roman', size=16, bold=False, color=None):
    run.font.name = latin
    run.font.size = Pt(size)
    run.bold = bold
    run._element.rPr.rFonts.set(qn('w:eastAsia'), east)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_para_format(p, first_indent=True, line=28, before=0, after=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    pf = p.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = Pt(line)
    if first_indent:
        pf.first_line_indent = Pt(32)
    else:
        pf.first_line_indent = None


def add_body(doc, text, bold_prefix=None, indent=True):
    p = doc.add_paragraph()
    set_para_format(p, first_indent=indent)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, east='FangSong', size=16, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, east='FangSong', size=16)
    else:
        r = p.add_run(text)
        set_run_font(r, east='FangSong', size=16)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.first_line_indent = None
    pf.space_before = Pt(8 if level == 1 else 4)
    pf.space_after = Pt(0)
    pf.line_spacing = Pt(28)
    pf.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, east='SimHei', size=16, bold=True)
    return p


def add_item(doc, text):
    p = doc.add_paragraph()
    set_para_format(p, first_indent=False)
    p.paragraph_format.left_indent = Pt(32)
    p.paragraph_format.first_line_indent = Pt(-32)
    r = p.add_run(text)
    set_run_font(r, east='FangSong', size=16)
    return p


def add_footer_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('— ')
    set_run_font(r, east='SimSun', size=10.5)
    fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = ' PAGE '
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
    r._r.append(fldChar1); r._r.append(instrText); r._r.append(fldChar2)
    r2 = p.add_run(' —')
    set_run_font(r2, east='SimSun', size=10.5)


def add_para_bottom_border(p, color=RED, size=18):
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.find(qn('w:pBdr'))
    if pBdr is None:
        pBdr = OxmlElement('w:pBdr')
        pPr.append(pBdr)
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(size))
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)


def setup_a4(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(3.0)
    sec.bottom_margin = Cm(2.7)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.6)
    sec.header_distance = Cm(1.5)
    sec.footer_distance = Cm(1.6)
    add_footer_page_number(sec)


def add_doc_styles(doc):
    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(16)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'FangSong')
    normal.paragraph_format.line_spacing = Pt(28)


def create_letter():
    doc = Document()
    setup_a4(doc)
    add_doc_styles(doc)

    # red head
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run('重庆市西部产教融合研究院文件')
    set_run_font(r, east='SimSun', size=28, bold=True, color=RED)
    add_para_bottom_border(p, RED, 20)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run('渝西产教研〔2026〕16号')
    set_run_font(r, east='FangSong', size=14)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(18)
    p.paragraph_format.line_spacing = Pt(32)
    title = '关于商请承办西部城市更新产教融合协同发展大会暨\n西部城市更新与钢结构智能建造专业建设联盟成立大会\n并共同推进首批建设任务的函'
    for i, part in enumerate(title.split('\n')):
        r = p.add_run(part)
        set_run_font(r, east='SimHei', size=22, bold=True)
        if i < 2:
            r.add_break()

    p = doc.add_paragraph()
    set_para_format(p, first_indent=False, align=WD_ALIGN_PARAGRAPH.LEFT)
    r = p.add_run('重庆建筑工程职业学院：')
    set_run_font(r, east='FangSong', size=16)

    paras = [
        '为深入贯彻国家关于城市更新、现代职业教育体系建设和产教融合有关部署，推动西部地区城市更新与钢结构智能建造领域产业需求、专业建设和人才培养有效衔接，重庆市西部产教融合研究院正会同有关行业企业、高等院校和科研机构筹备成立“西部城市更新与钢结构智能建造专业建设联盟”（以下简称“联盟”），拟于2026年8月25日在重庆召开西部城市更新产教融合协同发展大会暨联盟成立大会。',
        '经前期电话沟通，贵校对承办本次大会并参与联盟建设表达了积极意向。结合贵校在建筑工程、装配式建筑、智能建造及相关专业领域的建设基础、实践教学条件和校企合作资源，经研究，现商请贵校承办本次大会，并共同推进联盟首批专业建设和产教融合项目。',
        '本次合作不以单纯举办一次会议为目的，而是拟以大会为起点，以真实产业项目为载体，以专业建设和人才培养为主线，率先探索形成“学生有项目、教师有成果、企业有落地、学校有品牌”的产教融合建设模式。现将有关事项函商如下。'
    ]
    for t in paras: add_body(doc, t)

    add_heading(doc, '一、共同打造联盟首批样板建设项目')
    add_body(doc, '商请贵校在承办联盟成立大会的基础上，积极参与联盟首批建设，并结合学校现有专业、师资、实训基地和合作企业资源，率先开展真实产业项目进入教学、企业技术成果教育化转化和教师项目化成果建设。')
    add_body(doc, '研究院拟优先支持贵校申报联盟首批样板建设单位，通过一个阶段的建设，形成可展示、可评价、可复制、可推广的专业建设和产教融合成果，为联盟后续在西部地区有关院校推广提供实践样本。')

    add_heading(doc, '二、共同推动真实产业项目进入人才培养')
    add_body(doc, '围绕城市更新、钢结构智能建造、装配式建筑、BIM/CIM、智能装备、绿色建造、既有建筑改造等方向，双方共同从企业真实工程、产品、技术和应用场景中，择优遴选首批可进入学校教学的真实项目。')
    add_body(doc, '重点推动企业真实项目由“工程任务”转化为“教学任务”，形成可进入课程教学、综合实训、毕业设计、技能训练和学生创新实践的项目学习任务。')
    add_body(doc, '每个项目原则上应明确：')
    for s in ['（一）企业提供的真实项目、技术或案例；','（二）对应的专业、课程和教学环节；','（三）学生需要完成的具体项目任务；','（四）教师指导任务和教学成果要求；','（五）企业导师参与方式；','（六）学生最终形成的设计成果、技术方案、数字模型、项目报告、作品或其他能够证明专业能力的成果。']:
        add_item(doc, s)
    add_body(doc, '通过项目实施，让学生不仅“学过知识”，更能够形成真实参与项目、解决实际问题的能力证明。')

    add_heading(doc, '三、共同建立企业成果教育化转化机制')
    add_body(doc, '围绕企业进入学校过程中普遍存在的“产品进入不了教学、技术转化不了课程、案例形成不了任务、合作方案难以落地”等问题，双方共同探索企业成果教育化转化机制。')
    add_body(doc, '重点推动：')
    for s in ['企业产品向教学装备和实训条件转化；','企业技术向课程模块和教学内容转化；','工程案例向项目化教学任务转化；','技术方案向专业建设解决方案转化；','企业标准向岗位能力和人才培养标准转化；','真实工程项目向学生实践项目转化。']:
        add_item(doc, '（' + '一二三四五六'[['企业产品向教学装备和实训条件转化；','企业技术向课程模块和教学内容转化；','工程案例向项目化教学任务转化；','技术方案向专业建设解决方案转化；','企业标准向岗位能力和人才培养标准转化；','真实工程项目向学生实践项目转化。'].index(s)] + '）' + s)
    add_body(doc, '对具备条件的企业，由联盟组织学校专业教师、企业技术人员和有关专家共同开展需求分析和方案设计，形成学校能够理解、能够实施、能够评价的专业建设和校企合作项目。')
    add_body(doc, '原则上不以一般性框架协议和形式化签约作为主要成果，更加注重项目是否具有真实任务、实施主体、建设内容和成果产出。')

    add_heading(doc, '四、共同形成教师项目化专业发展成果')
    add_body(doc, '改变教师在传统校企合作中任务模糊、成果不清的问题。结合首批真实产业项目，建立教师“一项目、一任务、一课程、一成果”的项目化专业发展机制，推动教师在参与真实项目过程中形成：')
    teacher_items = ['（一）真实项目教学任务；','（二）项目化课程或课程模块；','（三）教学案例、教案及教学资源；','（四）企业实践和技术服务成果；','（五）教改课题、研究报告或典型案例；','（六）课程教材、活页式教材或数字教学资源；','（七）教学能力比赛、技能竞赛及其他教育教学成果的基础材料。']
    for s in teacher_items: add_item(doc, s)
    add_body(doc, '通过明确教师参与项目的职责、过程和成果，把校企合作由“配合企业”转变为教师专业能力提升和标志性成果培育的重要载体。')

    add_heading(doc, '五、共同形成大会首批实质性成果')
    add_body(doc, '本次大会拟突出“成立即建设、启动即落地”，原则上形成以下首批成果：')
    for s in ['（一）正式成立西部城市更新与钢结构智能建造专业建设联盟；','（二）发布联盟首批成员单位及首批样板建设单位；','（三）建立并发布联盟首批专家库；','（四）发布西部城市更新与钢结构智能建造有关产业人才需求和重点岗位需求成果；','（五）发布联盟首批“真实产业项目进教学”项目清单；','（六）围绕条件成熟的项目签署首批校企合作项目任务书或实质性合作协议。']:
        add_item(doc, s)
    add_body(doc, '对暂不具备实质合作条件的项目，不以大会现场签约为目的，待任务、责任主体和实施条件成熟后再行组织实施。')

    add_heading(doc, '六、拟请贵校重点支持事项')
    for s in ['（一）原则同意承办2026年8月25日西部城市更新产教融合协同发展大会暨联盟成立大会。','（二）明确1名校领导牵头，成立大会筹备和联盟建设工作组，并确定具体工作联系人。','（三）结合学校专业建设和产业合作基础，择优推荐5—10家城市更新、钢结构、装配式建筑、智能建造、数字技术、绿色建造等领域重点企业参与大会及后续项目建设。','（四）优先从中遴选2—3个具有真实技术、真实案例、真实工程任务和实际合作意愿的项目，作为联盟首批真实产业项目进入教学的候选项目。','（五）组织相关二级学院、专业负责人和骨干教师参与首批项目设计，并根据实际需要组建教师团队和学生项目团队。','（六）统筹提供大会主会场、现场观摩和必要的会务保障条件，展示贵校有关专业、实训基地和产教融合建设成果。','（七）与研究院共同研究贵校作为联盟首批样板建设单位的建设任务、阶段成果和后续推进机制。']:
        add_item(doc, s)

    add_heading(doc, '七、研究院重点承担事项')
    for s in ['（一）负责联盟成立大会总体方案、联盟建设方案和首批项目机制设计；','（二）负责协调联盟有关发起单位、高校、行业企业、专家及相关单位参会；','（三）负责联盟章程、成员单位、专家库、首批建设项目等核心材料的组织和确认；','（四）负责联盟标识、会议主视觉、会议手册、授牌、聘书、项目任务书等核心会议材料设计制作；','（五）组织有关专家参与企业真实项目向教学项目、专业建设方案和教师成果的教育化转化；','（六）负责大会核心议程、成果发布、项目对接、新闻宣传和有关组织协调工作；','（七）会后持续推动首批项目实施，组织阶段评价、成果总结和典型案例推广。']:
        add_item(doc, s)

    add_heading(doc, '八、工作衔接')
    add_body(doc, '鉴于大会筹备时间较紧，商请贵校于2026年8月12日前原则确认承办意向，并明确牵头校领导和工作联系人。')
    add_body(doc, '双方建立专项工作机制，共同研究大会方案、首批项目、参会企业、现场观摩及相关事项。具体会务安排和筹备任务按照双方确认的《西部城市更新产教融合协同发展大会暨西部城市更新与钢结构智能建造专业建设联盟成立大会会务安排及筹备任务清单》组织实施。')
    add_body(doc, '此次合作既是联盟成立大会的筹备工作，也是联盟首批专业建设项目的正式启动。希望双方充分发挥各自优势，共同探索形成真正能够进入专业、进入课程、进入项目、进入课堂的产教融合实践模式，为西部地区城市更新与智能建造领域人才培养和专业建设提供可复制、可推广的实践经验。')

    p = doc.add_paragraph()
    set_para_format(p, first_indent=True)
    r = p.add_run('专此函商，盼复。')
    set_run_font(r, east='FangSong', size=16)

    p = doc.add_paragraph()
    set_para_format(p, first_indent=False, align=WD_ALIGN_PARAGRAPH.LEFT)
    r = p.add_run('附件：西部城市更新产教融合协同发展大会暨西部城市更新与钢结构智能建造专业建设联盟成立大会会务安排及筹备任务清单')
    set_run_font(r, east='FangSong', size=16)

    # signature block
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run('重庆市西部产教融合研究院')
    set_run_font(r, east='FangSong', size=16)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run('2026年8月10日')
    set_run_font(r, east='FangSong', size=16)

    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(18)
    r = p.add_run('联系人：郭耀飞　　联系电话：19923115180')
    set_run_font(r, east='FangSong', size=14)

    path = OUT / '01_关于商请重庆建筑工程职业学院承办联盟成立大会并共同推进首批建设任务的函.docx'
    doc.save(path)
    return path


def style_table(table, col_widths=None, header=True, font_size=12.5):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    for i, row in enumerate(table.rows):
        set_no_split(row)
        if i == 0 and header:
            set_repeat_table_header(row)
        for j, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if i == 0 and header:
                set_cell_shading(cell, MID_GRAY)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (i == 0 or j == 0) else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = Pt(20)
                for run in p.runs:
                    set_run_font(run, east='SimSun', size=font_size, bold=(i==0))
    if col_widths:
        for row in table.rows:
            for idx, w in enumerate(col_widths):
                if idx < len(row.cells):
                    row.cells[idx].width = Cm(w)


def add_table(doc, headers, rows, col_widths=None, font_size=12.5):
    table = doc.add_table(rows=1, cols=len(headers))
    for j,h in enumerate(headers):
        table.rows[0].cells[j].text = h
    for row in rows:
        cells = table.add_row().cells
        for j,val in enumerate(row):
            cells[j].text = str(val)
    style_table(table, col_widths=col_widths, header=True, font_size=font_size)
    doc.add_paragraph()
    return table


def create_checklist():
    doc = Document()
    setup_a4(doc)
    add_doc_styles(doc)
    # narrower margins for tables
    sec = doc.sections[0]
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(6); p.paragraph_format.line_spacing = Pt(30)
    r = p.add_run('西部城市更新产教融合协同发展大会')
    set_run_font(r, east='SimHei', size=22, bold=True)
    r.add_break()
    r2 = p.add_run('暨西部城市更新与钢结构智能建造专业建设联盟成立大会')
    set_run_font(r2, east='SimHei', size=20, bold=True)
    r2.add_break()
    r3 = p.add_run('会务安排及筹备任务清单')
    set_run_font(r3, east='SimHei', size=22, bold=True)

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('（双方工作讨论稿）')
    set_run_font(r, east='KaiTi', size=14)

    add_body(doc, '为做好西部城市更新产教融合协同发展大会暨西部城市更新与钢结构智能建造专业建设联盟成立大会筹备工作，明确重庆市西部产教融合研究院与重庆建筑工程职业学院工作职责和关键时间节点，确保大会既完成联盟成立程序，又形成首批可落地建设项目，制定本工作清单。')

    add_heading(doc, '一、大会基本信息')
    rows = [
        ['会议名称','西部城市更新产教融合协同发展大会暨西部城市更新与钢结构智能建造专业建设联盟成立大会'],
        ['会议主题','城市更新赋能西部发展　产教融合共育智能建造人才'],
        ['会议时间','2026年8月25日（星期二）08:30—16:00'],
        ['拟定地点','重庆建筑工程职业学院'],
        ['会议规模','约80人，根据最终确认名单适当调整'],
        ['会议定位','坚持“成立即建设、启动即落地”，重点形成联盟首批成员、专家、项目和样板建设任务。']
    ]
    add_table(doc, ['项目','内容'], rows, col_widths=[3.2,12.8], font_size=12.5)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('大会建设逻辑：学生有项目、教师有成果、企业有落地、学校有品牌。')
    set_run_font(r, east='SimHei', size=15, bold=True)

    add_heading(doc, '二、大会拟形成的主要成果')
    for s in ['（一）完成联盟成立有关程序，正式启动联盟建设；','（二）发布联盟首批成员单位及首批样板建设单位；','（三）建立并发布联盟首批专家库；','（四）发布产业人才需求、重点岗位需求等首批研究成果；','（五）发布首批“真实产业项目进教学”项目清单；','（六）对条件成熟的2—3个项目签署项目任务书或实质性合作协议；','（七）确定重庆建筑工程职业学院首批样板建设有关任务；','（八）形成大会新闻宣传和首批典型项目传播材料。']:
        add_item(doc,s)

    add_heading(doc, '三、工作组织机制')
    add_heading(doc, '（一）成立联合筹备工作组', 2)
    add_body(doc, '由重庆市西部产教融合研究院、重庆建筑工程职业学院共同组建大会联合筹备工作组。建议重庆建筑工程职业学院明确1名校领导牵头，研究院明确1名负责人统筹，双方各指定1名具体联络人。')
    add_heading(doc, '（二）工作原则', 2)
    for s in ['1. 内容优先。先确定大会成果和项目，再安排会议形式。','2. 项目优先。签约必须有真实任务，不以框架协议数量作为大会成效。','3. 成果优先。所有首批项目均应明确学生任务、教师成果、企业投入和建设产出。','4. 安全有序。会务、考察、接待和后勤工作实行责任到人。','5. 信息确认。未最终确认的领导、单位、专家和企业，不提前进入正式印刷和公开宣传材料。']:
        add_item(doc,s)

    add_heading(doc, '四、重庆市西部产教融合研究院主要任务')
    sections = {
        '（一）大会总体设计':['1. 制定大会总体方案；','2. 设计大会核心议程；','3. 统筹联盟成立有关程序；','4. 统筹大会成果发布和项目启动安排。'],
        '（二）联盟建设材料':['1. 完善联盟章程及有关制度；','2. 汇总并确认首批联盟成员单位；','3. 确认首批样板建设单位；','4. 汇总并确认首批专家库成员；','5. 制作有关授牌、聘书及会议材料。'],
        '（三）外部资源协调':['1. 协调有关发起单位和重点合作单位；','2. 邀请有关教育行政部门、行业组织、高等院校、科研机构及专家代表；','3. 邀请中国建筑集团所属有关专业企业及城市更新、钢结构智能建造领域重点企业；','4. 统筹重要嘉宾接待需求。'],
        '（四）首批项目设计':['1. 与学校共同筛选真实产业项目；','2. 组织企业、学校教师和专家开展项目教育化转化；','3. 明确项目学生任务、教师任务、企业资源和预期成果；','4. 编制首批项目任务书；','5. 确认大会现场发布和签约项目。'],
        '（五）大会核心材料':['1. 大会主视觉设计；','2. 联盟LOGO及视觉应用；','3. 大会主PPT；','4. 主持词及核心串词；','5. 会议手册；','6. 联盟章程及审议材料；','7. 授牌和聘书；','8. 项目任务书及有关协议；','9. 成果发布材料；','10. 新闻通稿及宣传材料。']
    }
    for h, items in sections.items():
        add_heading(doc,h,2)
        for s in items: add_item(doc,s)

    add_heading(doc, '五、重庆建筑工程职业学院主要任务')
    school_sections = {
        '（一）成立工作组':['1. 明确1名校领导牵头；','2. 明确具体工作部门；','3. 确定1名日常联络人；','4. 组建会务、教学项目、企业联络和现场保障工作团队。'],
        '（二）落实会议场地':['1. 提供可满足约80人参会的主会场；','2. 配备LED屏、音响、话筒、演讲台、PPT播放等设备；','3. 安排签到区、贵宾休息区和必要的工作区域；','4. 做好会场布置、桌牌摆放和现场设备保障。'],
        '（三）组织企业资源':['结合学校现有合作企业，择优推荐5—10家与城市更新、钢结构设计制造施工、装配式建筑、BIM/CIM及建筑数字化、智能建造装备和建筑机器人、绿色建造和新型建筑材料、既有建筑检测改造和运维等方向相关的企业。','推荐企业重点考察其是否具有真实产品、真实技术、真实案例或者真实项目，而非单纯以参会数量为标准。'],
        '（四）筛选首批真实项目':['学校与研究院共同从推荐企业中筛选2—3个具备实际落地条件的首批项目。','每个候选项目应至少回答：企业能够提供什么真实资源；学校哪个专业或课程可以承接；学生具体完成什么任务；教师具体承担什么任务；最终形成什么教学或项目成果；项目由谁负责；多长时间能够完成。'],
        '（五）组织教师团队':['围绕首批项目，组织专业负责人、骨干教师和相关教学团队参加。','原则上每个项目明确学校项目负责人、企业技术负责人、参与教师、学生团队和项目成果要求。'],
        '（六）组织现场观摩':['结合学校实际条件，设计专业建设和产教融合成果观摩路线。','观摩重点突出学校现有专业建设基础、真实教学实训场景、企业技术在教学中的应用、学生项目成果以及下一阶段拟共同建设的真实项目。']
    }
    for h, items in school_sections.items():
        add_heading(doc,h,2)
        for s in items: add_body(doc,s)

    add_heading(doc, '六、双方共同推进事项')
    add_heading(doc, '（一）确认联盟首批建设成果', 2)
    for s in ['1. 首批成员单位名单；','2. 首批样板建设单位名单；','3. 首批专家名单；','4. 首批真实产业项目名单；','5. 首批项目签约或任务书清单。']:
        add_item(doc,s)
    add_heading(doc, '（二）共同完成项目教育化转化', 2)
    add_body(doc, '每个首批项目原则上形成一张《项目建设任务卡》，至少包括以下内容：')
    task_rows = [
        ['项目名称',''],['企业单位',''],['学校承接专业',''],['对应课程或教学环节',''],['企业提供资源',''],['学生学习任务',''],['教师建设任务',''],['企业导师职责',''],['主要成果',''],['项目负责人',''],['实施时间',''],['评价方式','']
    ]
    add_table(doc,['项目','填写内容'], task_rows, col_widths=[4.0,12.0], font_size=12)
    add_heading(doc, '（三）共同审核签约项目', 2)
    add_body(doc, '大会现场原则上只安排已经具备实质合作内容的项目，不得以增加签约数量为目的临时拼凑框架协议。项目至少应具备明确合作双方、明确实施任务、明确项目负责人、明确建设周期和明确阶段成果等基本条件。')

    add_heading(doc, '七、大会建议议程')
    agenda = [
        ['08:30—09:00','参会签到','参会代表签到、领取会议资料、引导就座。'],
        ['09:00—09:25','大会开幕','主持人介绍领导和嘉宾；承办学校领导致辞；有关发起单位或重要嘉宾致辞。'],
        ['09:25—10:05','联盟成立','介绍联盟筹备情况；审议联盟章程；宣布联盟成立；发布首批成员单位；首批单位代表授牌。'],
        ['10:05—10:35','专家与样板建设单位发布','发布联盟首批专家库；专家代表聘任；发布首批样板建设单位。'],
        ['10:35—11:20','主题交流','围绕城市更新、产业人才需求、钢结构智能建造和专业建设等安排主题报告或交流。'],
        ['11:20—12:00','首批项目发布','发布首批产业人才需求成果；发布首批真实产业项目进教学项目；对成熟项目签署项目任务书或实质性合作协议。'],
        ['12:00—13:30','工作午餐','工作午餐及交流。'],
        ['13:30—14:30','产教融合项目专题交流','重点交流企业产品如何转化为学校建设方案、企业技术如何进入课程、真实案例如何转化为学生项目、教师参与项目如何形成明确成果，并安排2—3个典型项目现场剖析。'],
        ['14:30—15:10','学校建设成果交流','由重庆建筑工程职业学院介绍相关专业、实训条件、校企合作和拟开展的样板项目。'],
        ['15:10—15:50','现场观摩','组织观摩有关实训基地、专业建设成果和真实教学场景；涉及生产或工程现场的，根据实际人数和现场安全要求准备必要防护用品。'],
        ['15:50—16:00','大会总结','明确联盟成立后的首批任务、项目负责人和下一阶段工作安排。']
    ]
    add_table(doc,['时间','环节','主要内容'],agenda,col_widths=[2.8,3.8,9.4],font_size=11.5)

    add_heading(doc, '八、参会人员建议构成')
    add_body(doc, '大会规模原则上控制在80人左右，主要包括有关教育行政部门、行业组织领导和代表，联盟有关发起单位代表，重庆市西部产教融合研究院有关人员，重庆建筑工程职业学院领导、二级学院及专业负责人，西部地区有关本科高校、职业本科、高职和中职学校代表，城市更新、钢结构、装配式建筑、智能建造等领域重点企业代表，行业、教育和工程技术专家，联盟首批成员单位及样板建设单位代表，以及有关媒体和工作人员。具体人员以最终确认名单为准。')

    add_heading(doc, '九、现场会务保障事项')
    support_sections = {
        '（一）签到及资料':['1. 设置签到区域；','2. 安排签到和引导人员；','3. 准备会议手册、参会证及相关资料；','4. 根据最终名单设置桌牌和座次。'],
        '（二）设备保障':['提前测试LED屏、音响、有线及无线话筒、演讲台、PPT播放电脑、翻页器以及备用设备和电源。'],
        '（三）餐饮保障':['根据最终确认参会人数安排工作午餐、饮用水及必要茶歇，提前统计特殊饮食需求。'],
        '（四）交通及停车':['明确停车区域、车辆引导和重要嘉宾车辆安排。'],
        '（五）医疗及安全':['1. 配备必要急救用品；','2. 明确现场安全责任人员；','3. 对现场观摩路线提前开展安全检查；','4. 涉及实训、生产或工程现场的，按要求准备安全防护用品。'],
        '（六）贵宾接待':['根据最终确认的领导和专家名单，安排必要的休息、接待和引导服务。']
    }
    for h, items in support_sections.items():
        add_heading(doc,h,2)
        for s in items: add_item(doc,s)

    add_heading(doc, '十、宣传和信息发布')
    for s in ['（一）大会新闻稿、正式名单、单位名称和人员职务发布前，由双方共同核实。','（二）未经最终确认的领导、单位、专家和企业不得提前进入正式印刷材料和公开宣传。','（三）大会宣传重点突出联盟为什么成立、解决什么真实问题、首批项目是什么、学生获得什么、教师形成什么成果、企业如何实现技术和项目落地、承办学校形成什么样板经验。','（四）避免将宣传重点停留在“授牌多少、签约多少、参会多少”等形式指标上。']:
        add_item(doc,s)

    add_heading(doc, '十一、关键时间节点')
    timeline = [
        ['8月12日前','原则确认承办；成立工作组；确定牵头校领导和联系人；初步确认会场；提交第一批推荐企业名单。','下发正式函件；提供大会总体方案；提供联盟筹备材料；启动有关单位和专家邀请。','建立专项工作机制。'],
        ['8月15日前','初步确认5—10家重点企业；提出2—3个真实产业项目；组织教师项目团队；初步确定现场观摩内容。','汇总参会单位和项目资源；组织项目初步论证。','初步确认参会单位、重点企业、首批项目和教师团队。'],
        ['8月18日前','确认会场、设备及现场观摩路线。','确认首批成员及样板建设单位；确认主要嘉宾和专家；完善首批项目内容。','确认大会正式议程。'],
        ['8月20日前','配合核定参会名单、项目人员和观摩安排。','完成联盟章程、项目任务书、会议主PPT和会议手册初稿；确认专家聘任和授牌名单。','基本确定最终参会名单。'],
        ['8月22日前','完成工作午餐及后勤安排；确认需要学校签署的项目文件。','完成授牌、聘书、新闻稿、主持词和会议资料印制。','完成核心材料联合审核。'],
        ['8月24日前','完成会场布置、设备调试、座次桌牌、签到资料和现场观摩安全检查。','核对全部会议物料。','完成全部会前检查。'],
        ['8月24日下午','配合全流程彩排。','组织全流程彩排。','重点演练开幕、联盟成立、授牌聘任、项目发布签约、PPT视频播放和现场观摩路线。'],
        ['8月25日','全程会务、教学项目、企业联络和现场保障。','会议统筹、主持、议程执行、成果发布和宣传。','按大会方案组织实施。']
    ]
    add_table(doc,['时间','重庆建筑工程职业学院','重庆市西部产教融合研究院','双方共同'],timeline,col_widths=[2.2,4.8,4.8,4.2],font_size=10.5)

    add_heading(doc, '十二、大会前必须完成的六项检查')
    checks = [
        ['一看联盟','章程、成员、专家和样板单位是否已经确认。'],
        ['二看项目','首批项目是否真实，是否有明确任务。'],
        ['三看学生','项目是否真正能够转化为学生学习任务。'],
        ['四看教师','参与教师最终形成什么成果是否明确。'],
        ['五看企业','企业提供的产品、技术、案例和项目如何进入学校是否明确。'],
        ['六看落地','每个重点项目是否有责任人、时间表和阶段成果。']
    ]
    add_table(doc,['检查项','检查标准'],checks,col_widths=[3.2,12.8],font_size=12)
    add_body(doc, '上述六项未明确的项目，原则上不作为大会重点签约和成果发布项目。')

    add_heading(doc, '十三、联络方式')
    contacts = [
        ['重庆市西部产教融合研究院','联系人：郭耀飞；联系电话：19923115180'],
        ['重庆建筑工程职业学院','牵头校领导：____________；工作部门：____________；联系人：____________；联系电话：____________；电子邮箱：____________']
    ]
    add_table(doc,['单位','联络信息'],contacts,col_widths=[5.2,10.8],font_size=12)

    for _ in range(2): doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run('重庆市西部产教融合研究院　　重庆建筑工程职业学院')
    set_run_font(r, east='FangSong', size=14)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run('2026年8月')
    set_run_font(r, east='FangSong', size=14)

    path = OUT / '02_联盟成立大会会务安排及筹备任务清单.docx'
    doc.save(path)
    return path


if __name__ == '__main__':
    p1 = create_letter()
    p2 = create_checklist()
    print(p1)
    print(p2)
