from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

BASE = Path(__file__).resolve().parents[1]
OUT = BASE / 'downloads'
OUT.mkdir(parents=True, exist_ok=True)

RED = 'C00000'
BLACK = '000000'
GRAY = '666666'
LIGHT = 'F3F3F3'


def set_run_font(run, east='SimSun', size=11, bold=False, color=BLACK):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.bold = bold
    run._element.rPr.rFonts.set(qn('w:eastAsia'), east)
    run.font.color.rgb = RGBColor.from_string(color)


def set_para(p, first=True, line=22, before=0, after=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    pf = p.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = Pt(line)
    pf.first_line_indent = Pt(22) if first else None


def add_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('— ')
    set_run_font(r, size=9, color=GRAY)
    fld1 = OxmlElement('w:fldChar'); fld1.set(qn('w:fldCharType'), 'begin')
    inst = OxmlElement('w:instrText'); inst.set(qn('xml:space'), 'preserve'); inst.text = ' PAGE '
    fld2 = OxmlElement('w:fldChar'); fld2.set(qn('w:fldCharType'), 'end')
    r._r.append(fld1); r._r.append(inst); r._r.append(fld2)
    r2 = p.add_run(' —')
    set_run_font(r2, size=9, color=GRAY)


def setup(doc, top=2.2, bottom=2.0, left=2.3, right=2.3):
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(top)
    sec.bottom_margin = Cm(bottom)
    sec.left_margin = Cm(left)
    sec.right_margin = Cm(right)
    sec.header_distance = Cm(1.2)
    sec.footer_distance = Cm(1.2)
    add_footer(sec)
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')


def add_title(doc, title, subtitle=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(title)
    set_run_font(r, east='SimHei', size=18, bold=True)
    if subtitle:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(8)
        r2 = p2.add_run(subtitle)
        set_run_font(r2, east='SimSun', size=10.5, color=GRAY)


def add_heading(doc, text):
    p = doc.add_paragraph()
    set_para(p, first=False, line=22, before=5, after=1, align=WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, east='SimHei', size=11.5, bold=True)
    return p


def add_clause(doc, text, red=False, first=True, bold=False):
    p = doc.add_paragraph()
    set_para(p, first=first)
    r = p.add_run(text)
    set_run_font(r, east='SimSun', size=11, bold=bold, color=RED if red else BLACK)
    return p


def add_mixed(doc, segments, first=True):
    p = doc.add_paragraph()
    set_para(p, first=first)
    for text, is_red, bold in segments:
        r = p.add_run(text)
        set_run_font(r, east='SimSun', size=11, bold=bold, color=RED if is_red else BLACK)
    return p


def shade(cell, fill=LIGHT):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def set_cell_text(cell, text, bold=False, center=False, color=BLACK):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_run_font(r, east='SimSun', size=10.5, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def party_table(doc):
    rows = [
        ('甲方', '招商局检测车辆技术研究院有限公司'),
        ('统一社会信用代码', '______________________________'),
        ('住所', '重庆市高新区新金大道9号'),
        ('法定代表人或授权代表', '______________________________'),
        ('乙方', '重庆市西部产教融合研究院'),
        ('统一社会信用代码', '______________________________'),
        ('住所', '重庆市两江新区红锦大道100号恒大中渝广场3号楼2210'),
        ('法定代表人或授权代表', '______________________________'),
    ]
    t = doc.add_table(rows=len(rows), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    t.autofit = False
    for i, (a, b) in enumerate(rows):
        t.columns[0].width = Cm(4.2)
        t.columns[1].width = Cm(11.6)
        set_cell_text(t.cell(i,0), a, bold=True, center=True)
        set_cell_text(t.cell(i,1), b)
        if i in (0,4):
            shade(t.cell(i,0)); shade(t.cell(i,1))
    return t


def create_agreement():
    doc = Document()
    setup(doc)
    add_title(doc, '智能网联新能源汽车测试技术产教融合项目', '合作框架协议（修订标记版）')
    party_table(doc)

    p = doc.add_paragraph()
    set_para(p, first=False, line=18, before=5, after=5, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run('说明：红色文字为本轮新增或修改内容；未标红内容原则沿用0807版本。')
    set_run_font(r, east='SimSun', size=9.5, bold=True, color=RED)

    add_clause(doc, '甲方具备智能网联新能源汽车研发测试、检测技术、设备平台、技术标准、工程师和产业资源；乙方具备院校合作拓展、产教融合研究、项目策划、教育教学体系设计、专家组织、课题研究和成果建设能力。双方本着合法合规、优势互补、权责一致、成果导向和长期共赢的原则，共同面向职业院校、应用型本科高校及其他教育机构，开发智能网联新能源汽车产业学院、实训基地、课程教材、师资培养、技术服务和标志性成果项目，依据《中华人民共和国民法典》及相关法律法规，订立本协议。', red=True)

    add_heading(doc, '第一条  协议性质与合作模式')
    add_clause(doc, '1.1 本协议为双方建立长期合作关系的框架性文件，明确合作方向、角色分工、项目开发、采购服务、成果交付、费用结算及风险边界。具体项目实行“一校一议、一项一签”，由双方在投标或发生实质投入前签署《具体学校项目合作确认书》及必要的专项合同。')
    add_clause(doc, '任何一方在具体项目确认文件签署前自行发生的调研、方案开发、样机、驻场人员、投标保证金、配套物料等成本原则上由投入方自行承担；双方另有书面约定，或因一方故意、重大过失或无正当理由退出已实质推进项目造成对方合理直接损失的除外。', red=True)
    add_clause(doc, '1.2 双方共同面向学校开展产业学院和产教融合项目交流洽谈，推动形成“甲方产业技术牵引、乙方教育方案与成果支撑、具体项目依法采购和实施”的合作模式。本协议不直接替代任何具体学校项目的投标文件、采购合同或设备合同，不构成学校必然立项、甲方必然中标或任何一方取得固定收益的承诺。')
    add_clause(doc, '双方为独立民事主体，本协议不构成合伙、联营或代理关系。未经另一方专项书面授权，任何一方不得代表另一方报价、签约、收款，或向院校作出超出另一方书面技术方案、服务方案、报价及实际交付能力的承诺；因本方提供资料错误、不完整或自身履约原因产生的责任，由本方依法承担。', red=True)

    add_heading(doc, '第二条  合作范围')
    add_clause(doc, '2.1 合作范围包括但不限于：（一）智能网联新能源汽车产业学院、实训中心及产教融合平台的策划、建设和运营；（二）专业群规划、人才培养方案、课程体系、教材和数字化教学资源开发；（三）教师技术培训、产业导师、双师型教师培养；（四）学生实训、实习就业和企业真实项目协同；（五）设备技术、测试方法、技术标准和产业项目导入；（六）课题、论文、白皮书、教学成果和品牌项目建设；（七）双方书面确认的其他产教融合产品和服务。')
    add_clause(doc, '2.2 超出前款合作范围且需对方投入资源、承担责任的新合作事项，应由双方另行书面确认。微信、邮件等可作为日常沟通和项目事实记录；涉及合作范围、费用、权利义务等实质变更的，以双方盖章、签字或依法有效的电子签署文件为准。', red=True)

    add_heading(doc, '第三条  甲方职责')
    add_clause(doc, '3.1 甲方作为项目产业技术牵头方，负责学校技术交流、设备与技术方案说明、技术参数、工程师资源、技术培训、产业资源和技术验收支持。')
    add_clause(doc, '甲方应根据具体项目需要，及时提供真实、合法、完整且具有实施基础的书面技术资料、技术参数、报价及交付条件，并在具体项目中按约承担设备供货、安装调试、技术培训、售后服务和技术验收等责任。因甲方提供资料不准确、不完整、未及时更新或履行不符合约定造成的责任，由甲方依法承担。', red=True)

    add_heading(doc, '第四条  乙方职责')
    add_clause(doc, '4.1 乙方负责院校合作拓展、项目策划和产业学院顶层设计，包括学校线索联系、需求调研、合作洽谈、项目提案、建设方案、专业与人才培养、课程教材、师资培养、学生实训和成果建设方案。', red=True)
    add_clause(doc, '乙方应在合理期限内向甲方真实同步已获取且与合作项目直接相关的院校需求、采购或招标信息；涉及乙方自身商业秘密、依法不得披露的信息除外。乙方制作涉及甲方技术的方案，应以甲方书面提供的最新技术资料为依据，不得擅自降低、篡改或曲解甲方技术指标。', red=True)

    add_heading(doc, '第五条  市场开发与项目确认')
    add_clause(doc, '5.1 双方原则上以共同合作项目的名义开展学校洽谈。任何一方对外使用另一方名称、品牌、人员、技术或成果，应事先取得确认。对已书面登记的项目，未经另一方书面同意，不得利用对方提供的学校关系、方案、报价、技术资料或成果绕开对方单独合作。项目合作保护期原则上为自书面登记之日起二十四个月。')
    add_clause(doc, '双方可根据客观事实将项目区分为一方存量项目、一方独立开发项目和双方共同开发项目。主张某院校或项目属于本协议签署前存量项目的，应能够提供此前已存在实质业务接洽的书面依据。任一方可通过双方约定邮箱或项目台账发起登记，另一方应在5个工作日内提出书面异议；逾期未提出异议，仅视为对登记事实无异议，不视为同意具体合作条件、报价或投入义务。', red=True)
    add_clause(doc, '保护期内已经进入立项、方案提交、采购申报、招投标、商务谈判等实质推进阶段的，项目保护可延续至该具体项目终止、落地或双方书面确认结束。', red=True)
    add_clause(doc, '5.2 每个学校项目在正式投标或发生重大投入前，双方应签署《具体学校项目合作确认书》，至少明确：学校与项目名称、采购方式、拟投标范围、双方工作任务、产品与成果清单、完成时间、验收标准、费用预算、知识产权、保密和项目联系人。')
    add_clause(doc, '具体项目确认文件还应根据实际情况明确签约主体、收款与开票主体、甲方设备及技术费用、乙方项目策划咨询/教育方案/成果建设等服务费用及支付节点、前期成本承担、流标废标处理、违约责任等事项。', red=True)

    add_heading(doc, '第六条  知识产权、保密与品牌使用')
    add_clause(doc, '6.1 双方合作前已经拥有的名称、品牌、技术、标准、资料和成果归原权利人所有。使用甲方技术资料开发课程、教材、数字资源和培训产品，应取得甲方专项书面授权。')
    add_clause(doc, '6.2 共同开发的课程、教材、报告、案例和成果，按照具体项目合同约定明确权属、使用权和署名。教学成果奖、课题、论文和案例申报应坚持真实贡献、规范署名和学术诚信。任何一方不得以对方名义开展与本项目无关的招生、收费、融资、担保或宣传活动。')
    add_clause(doc, '6.3 双方对合作过程中获悉的对方技术秘密、商业秘密、院校项目线索、需求调研数据、未公开方案及其他明确标识或依性质应当保密的信息承担对等保密义务。涉密资料的使用范围、留存、返还或销毁方式，由具体项目文件或专项保密协议约定；依法依规必须留存的档案资料除外。', red=True)
    add_clause(doc, '6.4 经权利方授权开展的课程、教材、数字资源或培训产品开发，其知识产权、使用范围及商业化安排按具体项目协议执行。未经授权，任何一方不得利用另一方专有技术、商业秘密、品牌、资质或其他知识产权开展对外商业活动；造成损失的，依法承担责任。', red=True)

    add_heading(doc, '第七条  期限、合规、违约责任与争议解决')
    add_clause(doc, '7.1 本协议有效期为三年，自双方签字盖章之日起生效。期满前六十日，双方可以协商续签。任一方需要提前终止本协议，应提前六十日书面通知另一方；重大违法违规或严重违约的，守约方有权立即暂停或解除，并依法获得相应赔偿。协议终止不影响已经中标、签约或正在履行项目的继续履行、结算、保密、知识产权、售后和责任承担。')
    add_clause(doc, '7.2 双方按照权责一致、过错与责任相适应原则承担违约责任。一般项目履约责任、违约金或责任上限由具体项目合同约定；未约定的，按照法律规定承担可预见的合理损失。对故意或重大过失、侵犯知识产权、泄露商业秘密、未经授权使用对方品牌或资质、商业贿赂等重大违法违约行为，依法承担相应责任。对能够补救且不影响合同目的实现的一般违约，守约方原则上应给予合理整改期限。', red=True)
    add_clause(doc, '7.3 双方在学校项目开发、采购、招投标、验收及费用结算过程中，应遵守法律法规和相关单位管理制度，不得通过商业贿赂、不正当利益输送等违法方式影响项目决策，不得以任何形式承诺学校必然立项、任何一方必然中标或规避依法应履行的采购程序。', red=True)
    add_clause(doc, '7.4 因本协议产生的争议，双方应先友好协商；协商不成的，任何一方可向被告住所地有管辖权的人民法院提起诉讼。', red=True)
    add_clause(doc, '7.5 本协议未尽事宜，由双方通过补充协议或具体项目协议另行约定。本协议一式肆份，双方各执贰份，具有同等法律效力。')

    doc.add_paragraph()
    t = doc.add_table(rows=5, cols=2)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    labels = [
        ('甲方（盖章）', '乙方（盖章）'),
        ('单位名称：招商局检测车辆技术研究院有限公司', '单位名称：重庆市西部产教融合研究院'),
        ('法定代表人或授权代表：', '法定代表人或授权代表：'),
        ('联系人及电话：', '联系人及电话：'),
        ('签署日期：____年__月__日', '签署日期：____年__月__日'),
    ]
    for i,(a,b) in enumerate(labels):
        set_cell_text(t.cell(i,0), a, bold=(i==0))
        set_cell_text(t.cell(i,1), b, bold=(i==0))

    path = OUT / '01_智能网联新能源汽车测试技术产教融合项目合作框架协议_修订标记版.docx'
    doc.save(path)
    return path


def create_reply():
    doc = Document()
    setup(doc, top=2.6, bottom=2.4, left=2.8, right=2.6)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run('重庆市西部产教融合研究院')
    set_run_font(r, east='SimSun', size=24, bold=True, color=RED)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '18'); bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), RED)
    pBdr.append(bottom); pPr.append(pBdr)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run('关于《智能网联新能源汽车测试技术产教融合项目合作协议》\n法务审核意见的回复函')
    set_run_font(r, east='SimHei', size=17, bold=True)

    p = doc.add_paragraph(); set_para(p, first=False, align=WD_ALIGN_PARAGRAPH.LEFT)
    r = p.add_run('招商局检测车辆技术研究院有限公司：')
    set_run_font(r, east='SimSun', size=11)

    paras = [
        '贵司法务团队对《智能网联新能源汽车测试技术产教融合项目合作协议》提出的审核意见已收悉。感谢贵司对协议文本和合作风险的认真审核。我院已结合双方前期沟通情况及后续项目实施需要，对相关意见进行了逐项研究。',
        '总体上，我院认同进一步明确合作边界、规范项目实施、加强技术资料和知识产权保护、防范双方履约风险的审核思路，也赞同继续采用“框架协议＋一校一议、一项一签”的合作方式。在不改变原协议精简框架的基础上，我院已形成修订标记版协议，红色文字为本轮新增或调整内容，供贵司法务直接对照审阅。',
    ]
    for t in paras: add_clause(doc, t)

    add_heading(doc, '一、关于本轮修订的基本原则')
    add_clause(doc, '本轮修订坚持“权责对等、边界清晰、框架精简、具体项目另签”的原则。对贵司提出的规范信息披露、技术资料保护、品牌及资质使用、条款编号等合理意见原则采纳；对涉及双方核心权益的事项，重点从责任对等和项目可执行角度作了平衡完善。')

    add_heading(doc, '二、重点完善事项')
    items = [
        '1. 双方职责对等。进一步明确甲方在技术资料、报价、设备交付、技术培训、售后及技术验收等方面的责任，同时明确乙方在院校合作拓展、需求对接、教育方案和成果建设方面的责任；未经专项书面授权，双方均不得代表对方作出承诺。',
        '2. 项目开发保护。保留24个月项目保护期，补充存量项目、独立开发项目和共同开发项目的客观认定及书面登记机制，并对保护期内已实质推进项目作延续安排。',
        '3. 费用与结算。一致坚持框架协议不承诺任何一方固定收益；具体学校项目再通过《具体学校项目合作确认书》或专项合同明确设备技术费用、项目策划咨询及教育成果服务费用、支付节点、成本承担和开票结算。',
        '4. 知识产权与双向保密。双方原有知识产权仍归原权利人；共同成果在具体项目中另行约定；技术资料、院校线索、需求数据、项目方案等均纳入对等保密范围。',
        '5. 违约与合规。按照权责一致、过错与责任相适应原则处理违约责任；一般履约责任由具体项目约定，重大违法违约依法承担责任。同时增加学校采购、招投标及廉洁合规要求。',
        '6. 争议解决。将原“协议签署地法院”调整为“被告住所地有管辖权的人民法院”，保持双方程序权利平衡。',
    ]
    for t in items: add_clause(doc, t, first=False)

    add_heading(doc, '三、后续建议')
    add_clause(doc, '建议双方以随函修订标记版为基础，由业务和法务对接人员直接对照修改。框架协议尽量保持精简，后续具体学校项目继续实行“一校一签”，将项目任务、采购方式、费用结算、成果交付和风险责任在具体项目文件中落实。')
    add_clause(doc, '我院重视与贵司的长期合作，也希望通过本轮法务审核把双方权利义务和合作边界进一步厘清，为后续产业学院及其他产教融合项目顺利推进奠定基础。')

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; p.paragraph_format.space_before = Pt(18)
    r = p.add_run('重庆市西部产教融合研究院')
    set_run_font(r, east='SimSun', size=11)
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r2 = p2.add_run('2026年8月18日')
    set_run_font(r2, east='SimSun', size=11)

    path = OUT / '02_关于合作协议法务审核意见的回复函.docx'
    doc.save(path)
    return path


if __name__ == '__main__':
    a = create_agreement()
    b = create_reply()
    print(a)
    print(b)
